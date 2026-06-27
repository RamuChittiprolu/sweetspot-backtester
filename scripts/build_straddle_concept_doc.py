from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("docs/Nifty_0920_Straddle_Basket_Concept_Explained.docx")


def set_run(run, bold=False, color="000000", size=11):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_fixed_table(table, widths):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")

            margins = tc_pr.find(qn("w:tcMar"))
            if margins is None:
                margins = OxmlElement("w:tcMar")
                tc_pr.append(margins)
            for name, value in [("top", 80), ("start", 120), ("bottom", 80), ("end", 120)]:
                node = margins.find(qn(f"w:{name}"))
                if node is None:
                    node = OxmlElement(f"w:{name}")
                    margins.append(node)
                node.set(qn("w:w"), str(value))
                node.set(qn("w:type"), "dxa")


def style_table(table, header=True):
    table.style = "Table Grid"
    if header:
        for cell in table.rows[0].cells:
            set_cell_shading(cell, "E8EEF5")
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    set_run(run, bold=True, color="0B2545", size=9.5)
    for row in table.rows[1 if header else 0 :]:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    set_run(run, size=9.5)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    r = p.add_run(text)
    set_run(r, bold=True, color="2E74B5" if level < 3 else "1F4D78", size=16 if level == 1 else 13 if level == 2 else 12)
    return p


def add_para(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(text)
    set_run(r)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(text)
    set_run(r)


def add_note(doc, title, lines, fill="F4F6F9", color="1F4D78"):
    table = doc.add_table(rows=1, cols=1)
    set_fixed_table(table, [9360])
    set_cell_shading(table.cell(0, 0), fill)
    p = table.cell(0, 0).paragraphs[0]
    r = p.add_run(title)
    set_run(r, bold=True, color=color)
    for line in lines:
        p = table.cell(0, 0).add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(line)
        set_run(r, size=10)
    doc.add_paragraph()


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    set_fixed_table(table, widths)
    for i, h in enumerate(headers):
        table.cell(0, i).text = h
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
    style_table(table)
    doc.add_paragraph()
    return table


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(4)
    r = title.add_run("NIFTY 09:20 Straddle Basket")
    set_run(r, bold=True, color="0B2545", size=22)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(12)
    r = subtitle.add_run("Concept Explained: What You Are Trading, Why P&L Differs, and How to Read the Workbook")
    set_run(r, color="1F4D78", size=12)

    add_note(
        doc,
        "Plain-English idea",
        [
            "This is a long-volatility strategy. You are not predicting whether NIFTY will go up or down.",
            "You are buying both CE and PE, so the basket benefits when option premiums expand enough after 09:20.",
            "The workbook's Daily Trades sheet shows the whole daily basket. The Leg Trades sheet shows each straddle component inside that basket.",
        ],
    )

    add_heading(doc, "1. What Is a Straddle?")
    add_para(doc, "A straddle means buying one Call option and one Put option at the same strike and same expiry.")
    add_table(
        doc,
        ["Part", "Meaning"],
        [
            ["CE", "Call option. Helps when market rises sharply or call premium expands."],
            ["PE", "Put option. Helps when market falls sharply or put premium expands."],
            ["Straddle", "CE + PE at the same strike."],
            ["Straddle premium", "CE premium + PE premium."],
        ],
        [2200, 7160],
    )

    add_heading(doc, "2. What Is This Basket?")
    add_para(doc, "This strategy buys two straddles at 09:20. That is why one daily trade contains two rows in the detailed sheet.")
    add_table(
        doc,
        ["Basket Component", "What It Contains", "Workbook Row"],
        [
            ["ATM -50 Straddle", "ATM -50 CE + ATM -50 PE", "One row in Leg Trades"],
            ["ATM Straddle", "ATM CE + ATM PE", "One row in Leg Trades"],
            ["Daily Basket", "ATM -50 Straddle + ATM Straddle", "One row in Daily Trades"],
        ],
        [2200, 4300, 2860],
    )

    add_note(
        doc,
        "Important",
        [
            "The Leg Trades sheet is named a little confusingly. Each row is not a single CE or PE option.",
            "Each row is one straddle component, meaning CE + PE combined for that strike.",
            "Daily Trades adds both straddle rows together for the date.",
        ],
        fill="FFF7ED",
        color="7C2D12",
    )

    add_heading(doc, "3. Entry Rules")
    for item in [
        "At 09:15, note the NIFTY spot/index open.",
        "Round the 09:15 open to the nearest 50-point strike. That is ATM.",
        "At 09:20, buy ATM -50 CE and ATM -50 PE.",
        "At 09:20, also buy ATM CE and ATM PE.",
        "Use the nearest available expiry on or after the trading date.",
        "For 130 quantity, every premium point equals Rs 130.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "4. Exit Rules")
    for item in [
        "Each straddle component has its own 15% target.",
        "Target premium = entry premium x 1.15.",
        "If the straddle target hits, that straddle exits at target.",
        "If target does not hit, exit that straddle at 15:15.",
        "The tested version has no intraday stop-loss, so losing days can be large.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "5. Why Daily P&L and Leg P&L Look Different")
    add_para(doc, "Daily Trades is the sum of the two straddle rows. Leg Trades shows only one straddle row at a time.")
    add_table(
        doc,
        ["Sheet", "Level", "What One Row Means"],
        [
            ["Leg Trades", "Straddle component", "Either ATM -50 straddle or ATM straddle"],
            ["Daily Trades", "Full basket", "Both straddles combined for that date"],
            ["Monthly", "Month", "All daily baskets added for the month"],
        ],
        [2100, 2500, 4760],
    )

    add_heading(doc, "6. Worked Example: 2025-08-29 Loss")
    add_para(doc, "This is the exact type of mismatch you saw in the screenshots. The Daily Trades row is larger because it includes two straddle rows.")
    add_table(
        doc,
        ["Component", "Entry", "Exit", "Points P&L", "Qty", "Gross Rs P&L"],
        [
            ["ATM -50 Straddle", "269.65", "183.25", "-86.40", "130", "-11,232"],
            ["ATM Straddle", "245.60", "176.60", "-69.00", "130", "-8,970"],
            ["Daily Basket Total", "515.25", "359.85", "-155.40", "130", "-20,202"],
        ],
        [2500, 1200, 1200, 1500, 900, 2060],
    )
    add_para(doc, "The daily gross rupee P&L is calculated like this:")
    add_table(
        doc,
        ["Formula", "Calculation", "Result"],
        [
            ["Points P&L", "-86.40 + -69.00", "-155.40 points"],
            ["Gross rupees", "-155.40 x 130", "-20,202"],
            ["Net rupees", "Gross rupees - charges", "Daily sheet net P&L"],
        ],
        [2500, 3400, 3460],
    )

    add_heading(doc, "7. Worked Example: 2026-03-04 Win")
    add_para(doc, "This shows how a strong volatility expansion day creates good P&L.")
    add_table(
        doc,
        ["Component", "Strike", "Entry", "Target/Exit", "Points P&L", "Gross Rs P&L"],
        [
            ["ATM -50 Straddle", "24350", "632.10", "726.91", "94.81", "12,325"],
            ["ATM Straddle", "24400", "625.45", "719.27", "93.82", "12,197"],
            ["Daily Basket Total", "-", "1257.55", "1446.18", "188.63", "24,522"],
        ],
        [2300, 1100, 1300, 1500, 1400, 1760],
    )

    add_heading(doc, "8. Why This Can Give Good P&L")
    for item in [
        "You are long both call and put premium, so a strong move in either direction can help.",
        "The target is only 15%, so the strategy tries to capture quick premium expansion.",
        "The basket has two nearby straddles, so it gets exposure around the ATM zone.",
        "High-premium or high-volatility days historically produced the strongest gains in the backtest.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "9. Main Risk")
    for item in [
        "If NIFTY stays quiet after 09:20, both CE and PE can decay.",
        "Because the tested version has no intraday stop-loss, losses can continue until 15:15.",
        "Charges reduce net P&L. The charges workbook already subtracts estimated brokerage, STT, stamp duty, exchange charge, SEBI fee, and GST.",
        "This is a research candidate, not a guaranteed income system.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "10. How to Read the Excel Workbook")
    add_table(
        doc,
        ["Sheet", "Use It For"],
        [
            ["Summary", "Overall assumptions, total gross P&L, charges, and net P&L."],
            ["Monthly", "Month-wise points, gross rupees, charges, net rupees, win rate, and capital used."],
            ["Daily Trades", "One row per date. This is the full basket result for the day."],
            ["Leg Trades", "One row per straddle component. Two rows usually make one daily basket."],
            ["Sources", "Charge assumptions and references."],
        ],
        [2200, 7160],
    )

    add_note(
        doc,
        "Mental model",
        [
            "Think of the strategy like buying two small volatility packages at 09:20.",
            "Package 1 is ATM -50 CE + PE.",
            "Package 2 is ATM CE + PE.",
            "Daily Trades shows both packages added together.",
        ],
    )

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    r = p.add_run("Source files: output/research_straddle and output/excel workbooks in this project.")
    set_run(r, color="555555", size=9)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)


if __name__ == "__main__":
    build()
