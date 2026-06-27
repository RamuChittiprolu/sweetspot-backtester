from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("docs/Nifty_0920_Straddle_Basket_Rules.docx")


def set_run(run, bold=False, color="000000", size=11):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(table, top=80, start=120, bottom=80, end=120):
    tbl_pr = table._tbl.tblPr
    margins = tbl_pr.find(qn("w:tblCellMar"))
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        tbl_pr.append(margins)
    for name, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_fixed_table(table, widths):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
    set_cell_margins(table)


def style_table(table, header=True):
    table.style = "Table Grid"
    if header:
        for cell in table.rows[0].cells:
            set_cell_shading(cell, "E8EEF5")
            for p in cell.paragraphs:
                for run in p.runs:
                    set_run(run, bold=True, color="0B2545", size=9.5)
    for row in table.rows[1 if header else 0 :]:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    set_run(run, size=9.5)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    r = p.add_run(text)
    set_run(r, bold=True, color="2E74B5" if level < 3 else "1F4D78", size=16 if level == 1 else 13 if level == 2 else 12)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(text)
    set_run(r)


def add_note_box(doc, title, lines):
    table = doc.add_table(rows=1, cols=1)
    set_fixed_table(table, [9360])
    set_cell_shading(table.cell(0, 0), "F4F6F9")
    p = table.cell(0, 0).paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run(r, bold=True, color="1F4D78")
    for line in lines:
        p = table.cell(0, 0).add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(line)
        set_run(r, size=10)
    doc.add_paragraph()


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    set_fixed_table(table, widths)
    for idx, header in enumerate(headers):
        table.cell(0, idx).text = header
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = str(value)
    style_table(table)
    doc.add_paragraph()
    return table


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(4)
    r = title.add_run("NIFTY 09:20 Long Straddle Basket")
    set_run(r, bold=True, color="0B2545", size=21)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(12)
    r = subtitle.add_run("Reference Rules, Backtest Summary, and Example Trades")
    set_run(r, color="1F4D78", size=13)

    add_note_box(
        doc,
        "Research status",
        [
            "This is a data-mined research candidate, not a finalized live system.",
            "It is the closest idea found so far to the 350-point/month target, but monthly results are uneven.",
            "The tested version has no intraday stop-loss; risk controls still need more work before live use.",
        ],
    )

    add_heading(doc, "Core Trade Rules")
    rules = [
        "At 09:15, read the NIFTY spot/index open.",
        "Round the 09:15 NIFTY open to the nearest 50-point strike. This is the ATM strike.",
        "At 09:20, enter two long straddles using the nearest available expiry on or after the trading date.",
        "Straddle 1: buy ATM -50 CE and ATM -50 PE.",
        "Straddle 2: buy ATM CE and ATM PE.",
        "For each straddle, combined entry premium = CE 09:20 close + PE 09:20 close.",
        "Target for each straddle = combined entry premium +15%.",
        "If target is hit intraday, exit that straddle at the target value.",
        "If target does not hit, exit at 15:15 using combined CE + PE close.",
        "The researched version uses no intraday stop-loss.",
    ]
    for rule in rules:
        add_bullet(doc, rule)

    add_heading(doc, "Formula")
    add_table(
        doc,
        ["Item", "Formula / Definition"],
        [
            ["ATM", "Nearest 50-point strike to NIFTY 09:15 open"],
            ["ATM -50 straddle", "(ATM -50 CE) + (ATM -50 PE)"],
            ["ATM straddle", "(ATM CE) + (ATM PE)"],
            ["Entry premium", "CE 09:20 close + PE 09:20 close"],
            ["Target", "Entry premium x 1.15"],
            ["Square-off", "15:15 combined close if target is not reached"],
        ],
        [2500, 6860],
    )

    add_heading(doc, "Backtest Summary")
    add_table(
        doc,
        ["Metric", "Value"],
        [
            ["Period used", "2025-09 to 2026-06; partial Aug 2025 excluded from averages"],
            ["Average monthly points", "343.68"],
            ["Best month", "1253.30"],
            ["Worst month", "-318.76"],
            ["Positive months", "9 out of 10"],
            ["Months above 350", "4 out of 10"],
            ["Best single leg", "09:20 ATM -50 straddle, +1907.72 total points, 62.83% win rate"],
            ["ATM leg", "09:20 ATM straddle, +1373.72 total points, 58.95% win rate"],
        ],
        [3100, 6260],
    )

    add_heading(doc, "Month-Wise Result")
    monthly_rows = [
        ["2025-09", "336.32"],
        ["2025-10", "472.30"],
        ["2025-11", "93.62"],
        ["2025-12", "-318.76"],
        ["2026-01", "412.97"],
        ["2026-02", "130.18"],
        ["2026-03", "1253.30"],
        ["2026-04", "489.70"],
        ["2026-05", "311.64"],
        ["2026-06", "255.56"],
    ]
    add_table(doc, ["Month", "Points"], monthly_rows, [2600, 6760])

    add_heading(doc, "Example Winning Days")
    add_heading(doc, "Example 1: 2026-03-04", 2)
    add_bullet(doc, "NIFTY 09:15 open: 24388.80; ATM: 24400.")
    add_table(
        doc,
        ["Leg", "Strike", "Entry", "Exit", "P&L", "Reason"],
        [
            ["ATM -50 straddle", "24350", "632.10", "726.91", "94.81", "TARGET"],
            ["ATM straddle", "24400", "625.45", "719.27", "93.82", "TARGET"],
            ["Total", "-", "1257.55", "1446.18", "188.63", "-"],
        ],
        [2100, 1200, 1400, 1400, 1300, 1960],
    )

    add_heading(doc, "Example 2: 2025-09-01", 2)
    add_bullet(doc, "NIFTY 09:15 open: 24432.70; ATM: 24450.")
    add_table(
        doc,
        ["Leg", "Strike", "Entry", "Exit", "P&L", "Reason"],
        [
            ["ATM -50 straddle", "24400", "199.90", "229.88", "29.98", "TARGET"],
            ["ATM straddle", "24450", "169.75", "195.21", "25.46", "TARGET"],
            ["Total", "-", "369.65", "425.10", "55.45", "-"],
        ],
        [2100, 1200, 1400, 1400, 1300, 1960],
    )

    add_heading(doc, "Example Losing Days")
    add_heading(doc, "Example 3: 2025-12-02", 2)
    add_bullet(doc, "NIFTY 09:15 open: 26087.95; ATM: 26100.")
    add_table(
        doc,
        ["Leg", "Strike", "Entry", "Exit", "P&L", "Reason"],
        [
            ["ATM -50 straddle", "26050", "139.00", "20.60", "-118.40", "SQUAREOFF"],
            ["ATM straddle", "26100", "114.00", "70.60", "-43.40", "SQUAREOFF"],
            ["Total", "-", "253.00", "91.20", "-161.80", "-"],
        ],
        [2100, 1200, 1400, 1400, 1300, 1960],
    )

    add_heading(doc, "Example 4: 2026-03-17", 2)
    add_bullet(doc, "NIFTY 09:15 open: 23493.20; ATM: 23500.")
    add_table(
        doc,
        ["Leg", "Strike", "Entry", "Exit", "P&L", "Reason"],
        [
            ["ATM -50 straddle", "23450", "181.80", "129.70", "-52.10", "SQUAREOFF"],
            ["ATM straddle", "23500", "206.60", "79.65", "-126.95", "SQUAREOFF"],
            ["Total", "-", "388.40", "209.35", "-179.05", "-"],
        ],
        [2100, 1200, 1400, 1400, 1300, 1960],
    )

    add_heading(doc, "Interpretation Notes")
    notes = [
        "This is a long-volatility idea. It benefits when combined CE + PE premium expands after 09:20.",
        "The strongest days tended to come when combined entry premium was already elevated.",
        "Low-volatility days can decay heavily because the researched version has no intraday SL.",
        "The average is close to the 350-point monthly target, but the result is not smooth enough to treat as stable yet.",
        "The next research step should test a volatility/premium filter and a daily maximum loss rule.",
    ]
    for note in notes:
        add_bullet(doc, note)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    r = p.add_run("Source: local backtest outputs in output/research_straddle.")
    set_run(r, color="555555", size=9)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)


if __name__ == "__main__":
    build()
