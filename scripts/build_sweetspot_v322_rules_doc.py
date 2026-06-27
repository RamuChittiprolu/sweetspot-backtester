from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "docs"
DOCX = OUT / "Sweet_Spot_v3_2_2_Reclaim_Continuation_Ignition_VIX_Rules.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(11, 37, 69)
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CAUTION = "FFF7ED"


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, bold: bool = False, fill: str | None = None) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(9.3)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.12
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if fill:
        shade_cell(cell, fill)


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "D9E2EC")


def set_table_widths(table, widths: list[float]) -> None:
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)
            tc_pr = row.cells[idx]._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def configure_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)


def bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)


def number(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)


def callout(doc: Document, title: str, body: str, fill: str = CAUTION) -> None:
    table = doc.add_table(rows=2, cols=1)
    set_table_borders(table)
    set_cell_text(table.cell(0, 0), title, bold=True, fill=LIGHT_BLUE)
    set_cell_text(table.cell(1, 0), body, fill=fill)
    table.cell(1, 0).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.add_paragraph()


def matrix(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    set_table_borders(table)
    set_table_widths(table, widths)
    for c, header in enumerate(headers):
        set_cell_text(table.cell(0, c), header, bold=True, fill=LIGHT_BLUE)
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row):
            set_cell_text(table.cell(r, c), value)
    doc.add_paragraph()


def label_table(doc: Document, rows: list[tuple[str, str]]) -> None:
    matrix(doc, ["Item", "Rule / Meaning"], [[a, b] for a, b in rows], [1.65, 4.85])


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(3)
    run = title.add_run("Sweet Spot v3.2.2 Reclaim + Continuation")
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = INK

    subtitle = doc.add_paragraph()
    subtitle.add_run("Detailed rules for first breakout, ignition reclaim, fresh reclaim, continuation, side switch, VIX, and exits").italic = True

    callout(
        doc,
        "Core Idea",
        "Do not trade every pivot touch. The strategy first keeps the original pivot breakout rules strict, then adds separate structure-based entries for reclaim and continuation moves that appear after the first noisy move has settled.",
        fill="F8FAFC",
    )

    heading(doc, "1. Strategy Identity")
    label_table(
        doc,
        [
            ("Strategy name", "Sweet Spot v3.2.2 Reclaim + Continuation"),
            ("Data", "5-minute NIFTY option candles plus NIFTY spot/index 09:15 open and India VIX 5-minute candles."),
            ("Strike", "ATM only in the current tested version."),
            ("Quantity model", "Backtest points first. Rupee P&L can be calculated later for chosen quantity."),
            ("Trade mode", "Only one active option trade at a time."),
            ("Sides", "CE and PE are scanned separately, but never traded together."),
        ],
    )

    heading(doc, "2. Pivot Calculation")
    number(doc, "At 09:15, read NIFTY spot/index open.")
    number(doc, "Select nearest ATM strike.")
    number(doc, "Read ATM CE 09:15 open and ATM PE 09:15 open.")
    number(doc, "Pivot = (CE 09:15 open + PE 09:15 open) / 2.")
    number(doc, "This pivot is the orange decision line for both CE and PE charts.")

    heading(doc, "3. Strict First Breakout Rules")
    matrix(
        doc,
        ["Filter", "Rule", "Purpose"],
        [
            ["Close above pivot", "Entry candle must close above pivot.", "Avoids buying below the decision line."],
            ["Entry ceiling", "Entry must be at or below pivot +20.", "Avoids chasing first breakout."],
            ["Range", "Candle range must be <=13 points.", "Keeps first breakout strict."],
            ["Body", "Body ratio must be >=0.35.", "Avoids weak candles."],
            ["Close quality", "Close must be in upper 60% of candle.", "Shows buyers held control."],
            ["Overextension", "Avoid if option already moved more than 35 points from day low.", "Avoids late first entries."],
            ["VIX", "VIX can hard-skip only if sharply falling and setup quality is weak.", "Avoids volatility crush without blocking strong structure."],
        ],
        [1.35, 2.7, 2.45],
    )

    heading(doc, "4. Ignition Reclaim Entry")
    callout(
        doc,
        "Why This Exists",
        "This was added because 2-Jun showed strong wide reclaim candles after a washout. The old anti-chase filter rejected them only because the candle range was large, even though the structure was strong.",
        fill="F8FAFC",
    )
    matrix(
        doc,
        ["Requirement", "Current Rule"],
        [
            ["Start time", "Only after 10:00 to avoid the noisy first 30-40 minutes."],
            ["Recent washout", "Recent candles must show a washout below pivot."],
            ["Reclaim", "Current candle must close back above pivot."],
            ["EMA", "Current candle should close above 20 EMA."],
            ["Entry ceiling", "Entry must be <= pivot +30."],
            ["Range", "Candle range allowed up to 70 points."],
            ["Body", "Body ratio must be >=0.60."],
            ["Close quality", "Close must be in upper 70% of candle."],
            ["Stop", "Exit on pivot close SL or base-low close SL."],
        ],
        [2.0, 4.5],
    )

    heading(doc, "5. Fresh Reclaim Entry")
    matrix(
        doc,
        ["Requirement", "Rule"],
        [
            ["Allowed after", "Previous trade exits by PIVOT_CLOSE_SL or C2C."],
            ["Wait", "At least 2 full candles after previous exit."],
            ["Same side", "Same side can re-enter only if price closes back above pivot."],
            ["Strength", "Reclaim candle must close above previous 2 candle highs."],
            ["EMA", "Reclaim candle should close above 20 EMA."],
            ["Entry ceiling", "Entry must be <= pivot +30."],
            ["Range", "Candle range <=18."],
            ["Stop", "Exit on close below pivot or close below reclaim base low."],
            ["Limit", "Max 3 trades per side per day in current test after ignition addition."],
        ],
        [2.0, 4.5],
    )

    heading(doc, "6. Continuation / Base Breakout Entry")
    matrix(
        doc,
        ["Requirement", "Rule"],
        [
            ["Context", "Price is already above pivot."],
            ["Base", "At least 3 candles holding above pivot or above 20 EMA."],
            ["Base high/low", "Base high and base low are calculated from those consolidation candles."],
            ["Trigger", "Enter when current candle closes above base high."],
            ["Entry ceiling", "Entry must be <= pivot +40."],
            ["Range", "Candle range <=18."],
            ["Overextension", "Do not use day-low overextension rule here."],
            ["Base distance", "Skip if entry is more than 20 points above recent base low."],
            ["Stop", "Exit if candle closes below base low."],
        ],
        [2.0, 4.5],
    )

    heading(doc, "7. Side Switch Entry")
    matrix(
        doc,
        ["Requirement", "Rule"],
        [
            ["No active trade", "Allowed only when there is no active trade."],
            ["Previous side", "Previous side must have exited by SL or C2C."],
            ["Opposite side strength", "Opposite side must close above pivot and 20 EMA."],
            ["Breakout", "Opposite side must close above previous 2 candle highs."],
            ["Entry ceiling", "Entry must be <= pivot +30."],
            ["Range", "Candle range <=18."],
            ["Limit", "Maximum 1 side switch per day."],
        ],
        [2.0, 4.5],
    )

    heading(doc, "8. VIX Rules")
    matrix(
        doc,
        ["VIX State", "How To Use It"],
        [
            ["OK", "Allow trades if price structure is valid."],
            ["MILD_FALL", "Do not automatically skip. Allow strong ignition, reclaim, and continuation structures."],
            ["SHARP_FALL", "Skip only when setup quality is weak. Strong structure may still be reviewed/allowed depending on rule settings."],
            ["REVIEW", "VIX missing or unmatched. Do not automatically skip; mark for review."],
            ["Important fix", "Use local exchange time when matching VIX. Do not convert +05:30 timestamps to UTC clock time."],
        ],
        [1.55, 4.95],
    )

    heading(doc, "9. Exit Rules")
    matrix(
        doc,
        ["Exit Type", "Rule"],
        [
            ["Pivot SL", "Always exit if candle closes below pivot, except continuation uses base-low stop as primary stop."],
            ["Base-low SL", "Ignition/reclaim/continuation setups use recent base low as additional structure stop."],
            ["C2C trigger", "After +20 points, move SL to cost-to-cost."],
            ["C2C exit", "Do not exit on intr candle wick touch. Exit C2C only if candle closes at or below entry after +20."],
            ["TSL trigger", "After +40 points, activate trailing stop using option candle lows."],
            ["Trailing SL", "Trail below candle lows and let the move continue."],
            ["Squareoff", "Exit open positions near end of day."],
        ],
        [1.55, 4.95],
    )

    heading(doc, "10. Real-Time Decision Flow")
    for step in [
        "Calculate ATM pivot at 09:15.",
        "Check whether CE or PE is above pivot. If both are below pivot, wait.",
        "For the first breakout, apply strict range/body/close/entry-ceiling filters.",
        "After 10:00, if there was a washout below pivot, watch for ignition reclaim.",
        "If an earlier trade exits, wait at least 2 full candles before considering fresh reclaim.",
        "If price is already above pivot and building a base, watch continuation/base breakout.",
        "Use VIX as a quality filter, not as a global trade killer.",
        "After entry, manage only one active trade until exit.",
        "Exit by pivot close SL, base-low SL, C2C close-back rule, trailing SL, or squareoff.",
    ]:
        number(doc, step)

    heading(doc, "11. June Examples")
    matrix(
        doc,
        ["Date", "What v3.2.2 Did", "Lesson"],
        [
            ["2-Jun", "Captured CE ignition at 10:15 for +33.20 and 12:30 for +130.95.", "Wide reclaim candles can be valid after washout; do not reject only because range is large."],
            ["3-Jun", "Captured the later CE rally structure instead of only the failed first pivot attempt.", "Fresh/continuation structure is more important than first touch."],
            ["5-Jun", "VIX now reads correctly as SHARP_FALL instead of REVIEW.", "Timezone bug fixed; VIX status should now be trusted except missing-date cases."],
            ["12-Jun", "VIX now reads MILD_FALL; trade still needs structure review.", "Mild VIX fall is not automatic skip."],
            ["16-Jun", "VIX now reads SHARP_FALL.", "Potential candidate for stricter quality skip."],
            ["19-Jun", "VIX now reads OK.", "No VIX excuse; evaluate setup and exit logic."],
            ["1-Jun", "Still REVIEW because the VIX CSV has no 2026-06-01 rows.", "Missing data should be marked review, not invented."],
        ],
        [1.0, 3.0, 2.5],
    )

    heading(doc, "12. Current June Backtest Snapshot")
    label_table(
        doc,
        [
            ("Report folder", "output/reclaim_continuation_v322/june_reports_vix_time_fixed"),
            ("v3.2.1 Structure v2", "-31.95 points"),
            ("v3.2.2 current", "+160.10 points"),
            ("Improvement", "+192.05 points"),
            ("Warning", "The month is improved, but continuation and some ignition trades still create losses. This is research, not a finished live system."),
        ],
    )

    callout(
        doc,
        "Bottom Line",
        "The current edge comes from separating setup types. First breakout must stay strict. Ignition reclaim can be wider only after a washout and after 10:00. VIX should explain trade quality and block only weak trades in sharply falling volatility.",
        fill="F8FAFC",
    )

    footer = doc.sections[0].footer.paragraphs[0]
    footer.text = "Sweet Spot v3.2.2 Reclaim + Continuation Rulebook"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.save(DOCX)
    print(DOCX)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
