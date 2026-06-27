from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("docs/Sweet_Spot_v3_2_1_Trade_Rules.docx")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(table, top=80, start=120, bottom=80, end=120):
    tbl_pr = table._tbl.tblPr
    margins = tbl_pr.find(qn("w:tblCellMar"))
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        tbl_pr.append(margins)
    for name, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_fixed_table(table, widths):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
    set_cell_margins(table)


def set_run(run, bold=False, color="000000", size=11):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_para(doc, text="", style=None, bold_prefix=None):
    p = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_run(r, bold=True)
        r = p.add_run(text[len(bold_prefix):])
        set_run(r)
    else:
        r = p.add_run(text)
        set_run(r)
    return p


def add_code(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_fixed_table(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F4F6F9")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("\n".join(lines))
    r.font.name = "Consolas"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor.from_string("1F4D78")


def add_rule(doc, number, title, bullets):
    h = doc.add_paragraph()
    h.style = doc.styles["Heading 2"]
    r = h.add_run(f"{number}. {title}")
    set_run(r, bold=True, color="2E74B5", size=13)
    for bullet in bullets:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.188)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        r = p.add_run(bullet)
        set_run(r)


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(4)
    r = title.add_run("Sweet Spot v3.2.1")
    set_run(r, bold=True, color="0B2545", size=22)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(14)
    r = subtitle.add_run("Multi-Strike Probability Backtester - Trading Rules")
    set_run(r, color="1F4D78", size=13)

    table = doc.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    set_fixed_table(table, [2700, 6660])
    rows = [
        ("Primary Data", "5-minute NIFTY option CSVs plus NIFTY spot/index candles."),
        ("Daily Anchor", "NIFTY 09:15 open selects the ATM strike."),
        ("Trade Limit", "Only one active option trade at a time."),
        ("Core Risk Rule", "Hard stop-loss is candle close below the same strike pivot."),
    ]
    for i, (label, detail) in enumerate(rows):
        table.cell(i, 0).text = label
        table.cell(i, 1).text = detail
        set_cell_shading(table.cell(i, 0), "E8EEF5")
        for cell in table.rows[i].cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    set_run(run, bold=(cell == table.cell(i, 0)))

    doc.add_paragraph()
    h1 = doc.add_paragraph(style="Heading 1")
    h1.add_run("Entry Framework")

    add_rule(doc, 1, "Daily ATM Selection", [
        "For each trading day, read the NIFTY spot/index candle at 09:15.",
        "Use the 09:15 NIFTY open to calculate the nearest 50-point ATM strike.",
        "Example: if NIFTY 09:15 open is 24,176, the nearest ATM is 24,200.",
    ])
    add_code(doc, ["ATM = nearest 50-point strike to NIFTY 09:15 open"])

    add_rule(doc, 2, "Strike Universe", [
        "Scan only the configured strike universe: ATM only, ATM +/-50, ATM +/-100, ATM +/-150, or all available strikes.",
        "The all-available mode may use only strikes that are available in real time; it must not choose strikes based on future day high/low.",
        "For each trading day and strike, use the nearest available expiry on or after the trading date.",
    ])

    add_rule(doc, 3, "Sweet Spot Pivot", [
        "For each selected strike, calculate one pivot from the same strike's 09:15 CE and PE open.",
        "The same pivot is used to evaluate both the CE and PE instruments for that strike.",
    ])
    add_code(doc, ["Pivot = (CE 09:15 Open + PE 09:15 Open) / 2"])

    add_rule(doc, 4, "No-Trade Condition", [
        "No trade is allowed if both CE and PE are below the pivot.",
        "A trade is considered only when CE or PE closes above the same strike pivot.",
    ])

    add_rule(doc, 5, "One Active Trade", [
        "CE and PE are scanned separately across the configured strikes.",
        "Only one instrument can be active at a time.",
        "When a trade is open, all new signals are ignored until that trade exits.",
    ])

    h1 = doc.add_paragraph(style="Heading 1")
    h1.add_run("Entry Filters and Setup Types")

    add_rule(doc, 6, "Entry Trigger", [
        "Entry is allowed only when the option candle closes above its pivot.",
        "Entry price is the signal candle close.",
        "Touches or wicks above the pivot are not enough; the candle must close above pivot.",
    ])
    add_code(doc, ["Close > Pivot"])

    add_rule(doc, 7, "Maximum Entry Price", [
        "Entry must be within pivot + 20 points.",
        "If the signal candle closes above pivot + 20, skip the trade.",
        "This prevents chasing stretched candles.",
    ])
    add_code(doc, ["Entry price <= Pivot + 20"])

    add_rule(doc, 8, "Entry Candle Range Filter", [
        "Calculate candle range as high minus low.",
        "Default maximum entry candle range is 13 points.",
        "Also test 15 points as a configurable variant.",
        "If the signal candle range is greater than the configured limit, skip the trade.",
    ])
    add_code(doc, ["Candle range = High - Low"])

    add_rule(doc, 9, "Retest/Bounce Near Pivot", [
        "Highest-priority setup.",
        "The candle must be green, close above pivot, stay within pivot + 20, and pass the entry candle range filter.",
        "The candle low must be near the pivot. Current default tolerance is low <= pivot + 3.",
    ])

    add_rule(doc, 10, "Vacuum Breakout", [
        "A strong green candle breaks from below and closes above pivot.",
        "The candle must close above the configured EMA.",
        "EMA period is configurable; the current test configs include EMA 15 and EMA 20.",
        "The entry must still be within pivot + 20 and pass the anti-chase candle range filter.",
        "Current body strength rule is body divided by candle range >= 0.55.",
    ])

    add_rule(doc, 11, "Normal Breakout", [
        "Previous close is below or equal to pivot.",
        "Current close is above pivot.",
        "Entry must remain within pivot + 20 and pass the entry candle range filter.",
    ])

    add_rule(doc, 12, "Same-Candle Signal Priority", [
        "If multiple valid signals appear on the same timestamp, choose only one.",
        "Priority order is retest/bounce, then vacuum breakout, then normal breakout.",
        "If still tied, choose the setup with the smaller distance from pivot.",
        "If still tied after that, choose smaller absolute strike distance, then option type as the final deterministic tie-break.",
    ])

    h1 = doc.add_paragraph(style="Heading 1")
    h1.add_run("Exit and Risk Management")

    add_rule(doc, 13, "Hard Stop-Loss", [
        "There is no fixed 15-point stop-loss.",
        "Exit when the active option candle closes below the same strike pivot.",
        "The exit price is the candle close that triggered the pivot-close stop.",
    ])
    add_code(doc, ["Exit when Close < Pivot"])

    add_rule(doc, 14, "Cost-to-Cost at +20", [
        "Once the trade reaches +20 points in favor, cost-to-cost protection activates.",
        "After C2C is active, if price returns to entry, exit exactly at entry price.",
        "The resulting P&L for a C2C exit is 0 points.",
    ])
    add_code(doc, ["C2C activates when High - Entry >= 20", "C2C exits when Low <= Entry Price"])

    add_rule(doc, 15, "Trailing Stop at +40", [
        "Once the trade reaches +40 points in favor, trailing stop-loss activates.",
        "The TSL trails option candle lows.",
        "The trailing stop only moves upward for long option trades.",
    ])
    add_code(doc, ["TSL = max(previous TSL, candle low)", "Exit when Low <= TSL"])

    add_rule(doc, 16, "C2C Remains Active After +40", [
        "C2C protection remains active even after trailing stop activates.",
        "After +20 has been reached, the trade should not become a loss under the C2C rule.",
    ])

    add_rule(doc, 17, "Square-Off and End-of-Data Exits", [
        "If no SL, C2C, or TSL exit happens before the configured square-off time, exit at the square-off candle close.",
        "Current default square-off time is 15:25.",
        "If there are no future candles after entry, the trade exits flat with reason NO_FUTURE_CANDLES.",
        "If data ends unexpectedly after entry, the trade exits with reason END_OF_DATA.",
    ])

    h1 = doc.add_paragraph(style="Heading 1")
    h1.add_run("Implementation Notes")
    add_rule(doc, 18, "Backtester-Specific Data Handling", [
        "Option CSVs are normalized to datetime, open, high, low, close, strike, option_type, and expiry_date.",
        "Strike, option type, and expiry can be extracted from filenames when columns are absent.",
        "Missing 09:15 spot candles, missing 09:15 option candles, missing CE/PE pivot pairs, duplicate candles, and candle gaps are reported as validation warnings.",
        "Duplicate candles for the same datetime, strike, option type, and expiry are de-duplicated by keeping the first occurrence.",
    ])

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    r = p.add_run("Generated for the Sweet Spot v3.2.1 Multi-Strike Probability Backtester project.")
    set_run(r, color="555555", size=9)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)


if __name__ == "__main__":
    build()
