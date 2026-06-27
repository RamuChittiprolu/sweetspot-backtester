from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "docs"
DOCX = OUT / "Sweet_Spot_v3_2_1_Normal_Pivot_Breakout_VIX_Structure_v2_Rules.docx"


BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(11, 37, 69)
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CAUTION = "FFF7ED"


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, bold: bool = False, fill: str | None = None) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(9.5)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.15
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if fill:
        shade_cell(cell, fill)


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "D9E2EC")


def set_table_widths(table, widths: list[float]) -> None:
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)
            tc_pr = row.cells[idx]._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    p.add_run(text)
    return p


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)


def add_number(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)


def add_callout(doc: Document, title: str, body: str, fill: str = CAUTION) -> None:
    table = doc.add_table(rows=2, cols=1)
    set_table_borders(table)
    set_cell_text(table.cell(0, 0), title, bold=True, fill=LIGHT_BLUE)
    set_cell_text(table.cell(1, 0), body, fill=fill)
    table.cell(1, 0).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.add_paragraph()


def add_label_table(doc: Document, rows: list[tuple[str, str]], widths: tuple[float, float] = (1.75, 4.75)) -> None:
    table = doc.add_table(rows=len(rows), cols=2)
    set_table_borders(table)
    set_table_widths(table, list(widths))
    for r, (label, detail) in enumerate(rows):
        set_cell_text(table.cell(r, 0), label, bold=True, fill=LIGHT_BLUE)
        set_cell_text(table.cell(r, 1), detail)
    doc.add_paragraph()


def add_matrix(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    set_table_borders(table)
    set_table_widths(table, widths)
    for c, header in enumerate(headers):
        set_cell_text(table.cell(0, c), header, bold=True, fill=LIGHT_BLUE)
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row):
            set_cell_text(table.cell(r, c), value)
    doc.add_paragraph()


def configure_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(3)
    run = title.add_run("Sweet Spot v3.2.1 Normal Pivot Breakout + VIX Structure v2")
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = INK

    subtitle = doc.add_paragraph()
    subtitle.add_run("Real-time trading rules, chart review notes, and June 2026 examples").italic = True

    add_callout(
        doc,
        "Purpose",
        "This rulebook is designed to trade only clean CE/PE pivot breakouts and retests. VIX is used only as a bad-trade avoidance filter, not as an entry signal. The goal is to reduce weak entries, avoid chasing, and avoid being shaken out too early after the first +20 move.",
        fill="F8FAFC",
    )

    add_heading(doc, "1. Strategy Identity", 1)
    add_label_table(
        doc,
        [
            ("Strategy name", "Sweet Spot v3.2.1 Normal Pivot Breakout + VIX Structure v2"),
            ("Instrument", "NIFTY weekly options, 5-minute candles"),
            ("Trade type", "Buy one option leg only: CE or PE. Do not buy both together."),
            ("Strike universe", "ATM only for this version. Later versions can test ATM +/-50 after ATM is stable."),
            ("Primary idea", "Use the 09:15 CE/PE pivot as the decision line. Trade only clean option premium strength above pivot."),
            ("VIX role", "Filter only. VIX should help avoid poor premium-expansion conditions; it should not create trades."),
        ],
    )

    add_heading(doc, "2. Pivot And Market Setup", 1)
    add_number(doc, "At 09:15, find the NIFTY spot/index open.")
    add_number(doc, "Select the nearest ATM strike from the 09:15 NIFTY open.")
    add_number(doc, "Read the 09:15 open of the selected strike CE and PE.")
    add_number(doc, "Calculate pivot: Pivot = (CE 09:15 open + PE 09:15 open) / 2.")
    add_number(doc, "Plot this pivot as the main orange decision line on both CE and PE charts.")

    add_callout(
        doc,
        "Example",
        "If ATM CE opens at 170 and ATM PE opens at 168 at 09:15, pivot = (170 + 168) / 2 = 169. Both CE and PE are judged against 169 separately.",
        fill="F8FAFC",
    )

    add_heading(doc, "3. Entry Rules", 1)
    add_matrix(
        doc,
        ["Rule", "Requirement", "Reason"],
        [
            ["One side only", "Trade CE or PE, never both at once.", "Keeps risk controlled and avoids mixed signals."],
            ["No-trade zone", "If both CE and PE are below pivot, wait.", "Premium is weak on both sides."],
            ["Valid close", "Entry candle must close above pivot.", "Avoids buying below the decision line."],
            ["Entry ceiling", "Entry must be within pivot +20.", "Prevents chasing stretched candles."],
            ["Candle range", "Entry candle range should be <=13 points.", "Filters late/impulsive candles."],
            ["Candle body", "Body ratio should be at least 0.35.", "Avoids indecision candles."],
            ["Close quality", "Close should be in the upper 60% of the candle.", "Shows buyers held control into close."],
            ["Overextension", "Avoid if option already moved more than about 35 points from day low before entry.", "Avoids late entries after the easy move is gone."],
        ],
        [1.35, 2.75, 2.4],
    )

    add_heading(doc, "4. Setup Types", 1)
    add_heading(doc, "A. Retest / Bounce Entry", 2)
    add_bullet(doc, "Price breaks above pivot, comes back near pivot, takes support, and closes back above pivot.")
    add_bullet(doc, "Entry candle should be green or clearly constructive, not a weak doji/indecision candle.")
    add_bullet(doc, "This has priority over normal breakout because risk is usually smaller and structure is clearer.")

    add_heading(doc, "B. Normal Pivot Breakout", 2)
    add_bullet(doc, "Previous candle was at or below pivot, and current candle closes above pivot.")
    add_bullet(doc, "Entry is valid only if all anti-chase filters pass: pivot +20 ceiling, range <=13, body quality, and close quality.")
    add_bullet(doc, "Do not treat a huge vertical candle as a clean breakout if the move is already extended.")

    add_heading(doc, "C. Fresh Reclaim After Failed Pivot Break", 2)
    add_bullet(doc, "If the first pivot breakout fails quickly, do not assume the day is over.")
    add_bullet(doc, "A fresh reclaim can be valid only if price rebuilds structure, closes back above pivot with strength, and does not look stretched.")
    add_bullet(doc, "This is the missing rule that explains why 3-Jun 13:10 looked better than the first 13:00 entry.")

    add_heading(doc, "5. VIX Avoidance Filter", 1)
    add_matrix(
        doc,
        ["VIX Check", "TAKE Condition", "Skip / Review Condition"],
        [
            ["VIX vs 09:15", "VIX at entry is not below 09:15 VIX open.", "Skip if VIX is below 09:15 open; premium may be compressing."],
            ["15-minute VIX trend", "VIX is flat or rising into entry.", "Skip if VIX is falling into entry."],
            ["Missing VIX", "Mark as REVIEW.", "Do not automatically trust the filter if VIX data is missing."],
        ],
        [1.7, 2.45, 2.35],
    )
    add_callout(
        doc,
        "Important",
        "VIX should reduce bad trades, not force trades. If price structure is poor, skip even when VIX says TAKE. If structure is excellent but VIX says SKIP, mark it for review rather than blindly accepting it.",
    )

    add_heading(doc, "6. Exit Rules", 1)
    add_matrix(
        doc,
        ["Exit Rule", "Action", "Notes"],
        [
            ["Hard SL", "Exit if a candle closes below pivot.", "This protects against false breakouts."],
            ["+20 rule", "Move SL to cost-to-cost after +20 points.", "Structure v2 waits for close back near entry instead of exiting on every wick touch."],
            ["+40 rule", "Activate trailing SL after +40 points.", "Trail below recent option candle lows."],
            ["Trailing method", "Use candle lows, not fixed 15-point SL.", "Let strong option trends continue."],
            ["Squareoff", "Exit near end of day if still open.", "Avoid overnight option risk."],
        ],
        [1.35, 2.55, 2.6],
    )

    add_heading(doc, "7. Simple Real-Time Decision Flow", 1)
    for step in [
        "Find ATM and calculate pivot at 09:15.",
        "Watch CE and PE separately; only one side can be traded.",
        "Ignore the trade if both sides are below pivot.",
        "When one side closes above pivot, check if price is within pivot +20.",
        "Check candle quality: range <=13, body ratio >=0.35, close in upper 60%.",
        "Check overextension: avoid if the option has already moved more than about 35 points from day low.",
        "Check VIX: avoid if VIX is below 09:15 open or falling into entry.",
        "Enter only if structure and filter both support the trade.",
        "Exit on pivot close SL, C2C close-back rule, trailing SL after +40, or squareoff.",
    ]:
        add_number(doc, step)

    add_heading(doc, "8. June 2026 Real-Time Examples", 1)
    add_matrix(
        doc,
        ["Date / Time", "What Happened", "Rule Lesson"],
        [
            ["1-Jun 11:50", "Trade reached +20 and old C2C exited at entry, then later rally continued.", "C2C should not be too jumpy. Structure v2 waits for a close back near entry instead of every wick touch."],
            ["3-Jun 13:00 / 13:10", "First CE reclaim failed quickly, but 13:10 showed a much cleaner rally structure.", "Add fresh reclaim rule after failed pivot break. Do not block the day after only one failed attempt if fresh structure appears."],
            ["4-Jun 13:30", "Old entry exited at C2C around 14:00 and missed later move.", "C2C protection is useful, but exit should respect candle close/structure, not only intr candle wick."],
            ["12-Jun", "Trade touched +20, exited C2C, then later CE rally developed around 13:30.", "A later side switch or fresh CE setup may be needed; do not assume first side's C2C means day is finished."],
            ["17-Jun 09:45", "CE had already rallied strongly before the entry and then failed.", "Avoid overextended entries. If easy move is already 35-40 points old, wait for a new base."],
            ["22-Jun 09:20", "Entry candle was indecisive and failed.", "Require candle body and close quality. Avoid doji/weak candles even when pivot condition technically passes."],
        ],
        [1.25, 2.8, 2.45],
    )

    add_heading(doc, "9. Do / Do Not Checklist", 1)
    add_matrix(
        doc,
        ["Do", "Do Not"],
        [
            ["Wait for a clean close above pivot.", "Do not buy just because price touched above pivot."],
            ["Prefer retest/bounce near pivot.", "Do not chase candles far above pivot."],
            ["Use VIX to avoid weak premium conditions.", "Do not use VIX as the entry trigger."],
            ["Avoid overextended first-move entries.", "Do not enter after a 40-point rally just because rules technically pass."],
            ["Allow fresh structure after a failed first entry.", "Do not repeatedly re-enter the same weak chop."],
            ["Review chart context before live execution.", "Do not blindly trust one backtest row without visual confirmation."],
        ],
        [3.25, 3.25],
    )

    add_heading(doc, "10. Current Status From Backtest", 1)
    add_label_table(
        doc,
        [
            ("June result", "Structure v2 improved June from -44.00 points to -31.95 points, but it is still not profitable."),
            ("VIX-only takeaway", "VIX helped skip some trades, but did not create a positive edge by itself."),
            ("Main next improvement", "The next useful rule is fresh structure / second breakout after a failed pivot attempt."),
            ("Live-trading caution", "This is a research rulebook, not yet a proven live strategy. More months must be tested before using real capital."),
        ],
    )

    add_callout(
        doc,
        "Bottom Line",
        "The current edge is not in simply taking every pivot break. The edge, if it exists, is likely in clean reclaim structure after failed moves, avoiding overextended entries, and using VIX only to reject bad environments.",
        fill="F8FAFC",
    )

    footer = doc.sections[0].footer.paragraphs[0]
    footer.text = "Sweet Spot v3.2.1 Structure v2 Rulebook"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.save(DOCX)
    print(DOCX)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
